# -*- coding: utf-8 -*-
"""office_utils —— Office 文档(docx/xlsx)读写工具包

从一个练习 demo 逐步封装而来: 把"碰巧能用"的代码提炼成"知道边界"的工具函数,
目标场景是"射频组件测试结果 → word 报告 / excel 记录"。

设计原则:
- 按功能域分模块, 不按读/写分
- __init__.py 统一导出, 调用方永远不感知内部文件结构
- 工具函数无状态, 不封装类 (类/配置留给将来高层 API)

使用示例:
    from docx import Document
    from office_utils import fill_bookmarks, set_cell, add_rows, set_doc_fonts

    doc = Document("template.docx")
    set_doc_fonts(doc, "微软雅黑")                  # 一次性修复中文字体
    fill_bookmarks(doc, {...}, font="微软雅黑")     # 填书签
    table = doc.tables[0]
    set_cell(table.cell(1, 0), "项目名", bold=True) # 填表
    add_rows(table, [["A", 1.5], ["B", 2.3]])       # 动态加数据
    doc.save("output.docx")
"""
from .docx_font import set_run_font, set_style_font
from .docx_common import set_doc_fonts, DEFAULT_STYLE_NAMES
from .bookmark import (
    add_bookmark,
    find_bookmark, find_bookmarks,
    fill_bookmark, fill_bookmarks,
    get_bookmark_text,
)
from .table import (
    fill_table,
    find_table_by_header, find_cell_by_text,
    set_cell, format_cell, format_row,
    add_rows,
    set_repeat_header,
)


__all__ = [
    # 字体
    "set_run_font", "set_style_font",
    # 文档级
    "set_doc_fonts", "DEFAULT_STYLE_NAMES",
    # 书签
    "add_bookmark",
    "find_bookmark", "find_bookmarks",
    "fill_bookmark", "fill_bookmarks",
    "get_bookmark_text",
    # 表格
    "fill_table",
    "find_table_by_header", "find_cell_by_text",
    "set_cell", "format_cell", "format_row",
    "add_rows",
    "set_repeat_header",
]
